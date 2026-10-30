"""Generate scope-isolated ESI queue/audit reports without raw source payloads."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _mask(value: Any) -> str:
    raw = "" if value is None else str(value)
    return ("*" * max(0, len(raw) - 4) + raw[-4:]) if raw else ""


def _public_row(row: dict[str, Any]) -> dict[str, Any]:
    history = row.get("status_history") if isinstance(row.get("status_history"), list) else []
    return {
        "Record ID": row.get("record_id", ""),
        "Company": row.get("company", ""),
        "Deputed Company": row.get("deputed_company", ""),
        "Source": row.get("source", ""),
        "Employee Code": row.get("employee_code", ""),
        "Employee Name": row.get("name", ""),
        "Aadhaar (masked)": _mask(row.get("aadhaar", "")),
        "DOB": row.get("dob", ""),
        "DOJ": row.get("doj", ""),
        "Gender": row.get("gender", ""),
        "Mobile": _mask(row.get("mobile", "")),
        "Present State": row.get("present_state", ""),
        "Present District": row.get("present_district", ""),
        "Bank IFSC": row.get("bank_ifsc", ""),
        "Status": row.get("status", ""),
        "Attempts": row.get("attempt_count", 0),
        "Retries": row.get("retry_count", 0),
        "Duplicates": row.get("duplicate_count", 0),
        "Last Message": row.get("last_message", ""),
        "First Seen": row.get("first_seen_at", ""),
        "Last Seen": row.get("last_seen_at", ""),
        "Last Success": row.get("last_success_at", ""),
        "Last Failure": row.get("last_failure_at", ""),
        "History Events": len(history),
    }


def build_esi_report(
    rows: list[dict[str, Any]],
    events: list[dict[str, Any]],
    *,
    company_id: str,
    deputed_company_id: str,
    company_name: str = "",
    deputed_company_name: str = "",
) -> dict[str, Any]:
    public_rows = [_public_row(row) for row in rows]
    statuses = [str(row.get("Status", "")) for row in public_rows]
    summary = {
        "module": "ESI",
        "company_id": str(company_id),
        "deputed_company_id": str(deputed_company_id),
        "company": company_name or (public_rows[0].get("Company", "") if public_rows else ""),
        "deputed_company": deputed_company_name or (public_rows[0].get("Deputed Company", "") if public_rows else ""),
        "total": len(public_rows),
        "success": sum(status in {"ESI_COMPLETED", "ESI_REGISTERED"} for status in statuses),
        "failed": sum("FAILED" in status or "VALIDATION_FAILED" in status for status in statuses),
        "pending": sum(status not in {"ESI_COMPLETED", "ESI_REGISTERED"} and "FAILED" not in status and "VALIDATION_FAILED" not in status for status in statuses),
        "audit_events": len(events),
        "generated_at": _now(),
    }
    return {"summary": summary, "records": public_rows, "events": events}


def workbook_bytes(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    for key, value in report.get("summary", {}).items():
        summary.append([key, value])
    records = workbook.create_sheet("Records")
    rows = report.get("records", [])
    headers = list(rows[0].keys()) if rows else ["Record ID", "Company", "Deputed Company", "Status", "Last Message"]
    records.append(headers)
    for row in rows:
        records.append([row.get(header, "") for header in headers])
    events = workbook.create_sheet("Audit Events")
    events.append(["Timestamp", "Event", "Status", "Record ID", "Message"])
    for item in report.get("events", []):
        events.append([item.get("timestamp", ""), item.get("event", ""), item.get("status", ""), item.get("record_id", ""), item.get("message", "")])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max((len(str(cell.value or "")) for cell in column), default=10)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), 55)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def json_bytes(report: dict[str, Any]) -> bytes:
    return json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8")

