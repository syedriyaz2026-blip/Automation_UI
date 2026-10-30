from pathlib import Path
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import PatternFill


headers = [
    "Sl. No", "Emp Code", "Company", "Name", "DOJ", "DOB", "Aadhar Number",
    "Father Name", "Marital Status", "Husband/Wife Name", "Gender", "Mobile",
    "Location", "Basic Wage", "UAN No"
]

rows = [
    [1, "DEMO001", "Example Consulting Ltd", "Demo Employee One", "05/10/2026", "19/09/1997", "762330210052", "Demo Father", "Un-Married", "N/A", "Female", "9000000001", "Bengaluru", 18600, "100000000001"],
    [2, "DEMO002", "Example Consulting Ltd", "Demo Employee Two", "05/10/2026", "22/02/1995", "719680799020", "Demo Parent", "Married", "Demo Spouse", "Male", "9000000002", "Bengaluru", 22000, "100000000002"],
]

ROOT = Path(__file__).resolve().parents[4]
WORKBOOK_DIR = ROOT / "PF" / "backend" / "samples" / "workbooks"
WORKBOOK_DIR.mkdir(parents=True, exist_ok=True)

path = WORKBOOK_DIR / "sample_employee_data.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Employees"
ws.append(headers)
for row in rows:
    ws.append(row)
ws.freeze_panes = "A2"
for cell in ws[1]:
    font = copy(cell.font)
    font.bold = True
    cell.font = font
for column in ws.columns:
    letter = column[0].column_letter
    ws.column_dimensions[letter].width = max(14, min(max(len(str(c.value or "")) for c in column) + 2, 30))
wb.save(path)
print(path)

mandatory_path = WORKBOOK_DIR / "PF_mandatory_template.xlsx"
mandatory_wb = Workbook()
mandatory_ws = mandatory_wb.active
mandatory_ws.title = "Mandatory Employees"
mandatory_headers = ["UAN No", "Aadhar Number", "Name", "DOB", "DOJ", "Basic Wage"]
mandatory_ws.append(mandatory_headers)
mandatory_ws.append(["", "", "", "", "", ""])
mandatory_ws.freeze_panes = "A2"
for cell in mandatory_ws[1]:
    font = copy(cell.font)
    font.bold = True
    cell.font = font
    cell.fill = PatternFill(fill_type="solid", fgColor="E8E5FE")
for column in mandatory_ws.columns:
    letter = column[0].column_letter
    mandatory_ws.column_dimensions[letter].width = max(16, len(str(column[0].value)) + 4)
mandatory_wb.save(mandatory_path)
print(mandatory_path)

bad_path = WORKBOOK_DIR / "sample_validation_cases.xlsx"
bad_wb = Workbook()
bad_ws = bad_wb.active
bad_ws.title = "Validation Cases"
bad_ws.append(headers)
bad_ws.append(rows[0])
bad_ws.append([3, "", "", "Bad@Name", "31/02/2026", "not-a-date", "123456789012", "", "Complicated", "", "Unknown", "12345", "", -100, "123"])
bad_ws.append([4, "DUP001", "Example Consulting Ltd", "Duplicate Employee", "06/10/2026", "20/09/1997", rows[0][6], "Demo Father", "Un-Married", "", "Female", "9000000003", "Bengaluru", 18000, rows[0][14]])
bad_ws.freeze_panes = "A2"
for cell in bad_ws[1]:
    font = copy(cell.font)
    font.bold = True
    cell.font = font
for column in bad_ws.columns:
    letter = column[0].column_letter
    bad_ws.column_dimensions[letter].width = max(14, min(max(len(str(c.value or "")) for c in column) + 2, 30))
bad_wb.save(bad_path)
print(bad_path)
