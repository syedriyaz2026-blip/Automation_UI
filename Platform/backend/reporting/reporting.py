from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook, load_workbook

HEADERS = [
    "Source Row", "Employee Code", "Employee Name", "UAN", "Status",
    "Error Type", "Workflow Step", "Message", "Started At", "Completed At",
    "Screenshot/Evidence"
]

def completed_rows(report_path: Path) -> set[str]:
    if not report_path.exists():
        return set()
    try:
        wb = load_workbook(report_path, read_only=True, data_only=True)
        values = list(wb.active.iter_rows(min_row=2, values_only=True))
        wb.close()
        return {str(row[0]) for row in values if row and row[0] is not None and row[4] == "Registered - Pending Approval"}
    except Exception:
        return set()


def uncertain_rows(report_path: Path) -> set[str]:
    """Rows whose Save result was not confirmed; never blindly resubmit them."""
    if not report_path.exists():
        return set()
    try:
        wb = load_workbook(report_path, read_only=True, data_only=True)
        values = list(wb.active.iter_rows(min_row=2, values_only=True))
        wb.close()
        return {str(row[0]) for row in values if row and row[0] is not None and row[4] == "Save Result Unknown"}
    except Exception:
        return set()
def existing_results(report_path: Path) -> list[dict[str, str]]:
    if not report_path.exists():
        return []
    try:
        wb = load_workbook(report_path, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(min_row=2, values_only=True))
        wb.close()
        return [dict(zip(HEADERS, ["" if value is None else str(value) for value in row])) for row in rows]
    except Exception:
        return []

def write_report(path: Path, results: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Run Log"
    ws.append(HEADERS)
    for result in results:
        ws.append([result.get(header, "") for header in HEADERS])
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 12), 55)
    ws.freeze_panes = "A2"
    wb.save(path)
