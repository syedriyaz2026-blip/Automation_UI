"""Generate a scope-isolated, downloadable report for a processing run."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook


SUCCESS_STATUS = "Registered - Pending Approval"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _mask_aadhaar(value: Any) -> str:
    raw = "" if value is None else str(value)
    digits = "".join(character for character in raw if character.isdigit())
    return ("*" * max(0, len(digits) - 4) + digits[-4:]) if digits else ""


def _safe_json(value: Any) -> Any:
    try:
        json.dumps(value)
        return value
    except (TypeError, ValueError):
        return str(value)


def _rows(run: dict[str, Any]) -> list[dict[str, Any]]:
    records = {str(item.get("record_id", "")): item for item in run.get("records", [])}
    output: list[dict[str, Any]] = []
    for result in run.get("results", []):
        record = records.get(str(result.get("record_id", "")), {})
        output.append({
            "Company": record.get("company", run.get("options", {}).get("company_name", "")),
            "Deputed Company": record.get("deputed_company", run.get("options", {}).get("deputed_company_name", "")),
            "Source": record.get("source", run.get("options", {}).get("source", "")),
            "Source Row": result.get("row", record.get("row_number", "")),
            "Employee Code": result.get("employee_code", record.get("emp_code", "")),
            "Employee Name": result.get("employee_name", record.get("name", "")),
            "UAN": result.get("uan", record.get("uan", "")),
            "Aadhaar (masked)": _mask_aadhaar(record.get("aadhaar", "")),
            "DOB": record.get("dob", ""),
            "DOJ": record.get("doj", ""),
            "Basic Wage": record.get("basic_wage", ""),
            "Status": result.get("status", ""),
            "Error Type": result.get("error_type", ""),
            "Workflow Step": result.get("workflow_step", ""),
            "Message": result.get("message", ""),
            "Evidence": result.get("screenshot", ""),
            "Completed At": result.get("completed_at", ""),
        })
    return output


def generate_run_report(run: dict[str, Any]) -> dict[str, Any]:
    """Write XLSX and JSON reports atomically and return public metadata.

    Credentials and raw source payloads are intentionally excluded. Aadhaar is
    masked in the downloadable report while the local evidence remains linked
    by its path/name.
    """
    run_id = str(run["run_id"])
    root = Path(__file__).resolve().parents[1] / "api_runs" / run_id
    root.mkdir(parents=True, exist_ok=True)
    rows = _rows(run)
    summary = {
        "run_id": run_id,
        "company_id": run.get("options", {}).get("company_id", ""),
        "deputed_company_id": run.get("options", {}).get("deputed_company_id", ""),
        "status": run.get("status", ""),
        "total": run.get("total", len(rows)),
        "completed": run.get("completed", 0),
        "success_count": sum(item.get("Status") == SUCCESS_STATUS for item in rows),
        "failed_count": sum(item.get("Status") not in {"", SUCCESS_STATUS} for item in rows),
        "generated_at": _now(),
    }
    json_payload = {
        "summary": summary,
        "results": rows,
        "events": [_safe_json(item) for item in run.get("events", [])],
    }
    json_path = root / "processing_report.json"
    json_tmp = json_path.with_suffix(".json.tmp")
    json_tmp.write_text(json.dumps(json_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    json_tmp.replace(json_path)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    for key, value in summary.items():
        summary_sheet.append([key, value])
    results_sheet = workbook.create_sheet("Results")
    headers = list(rows[0].keys()) if rows else [
        "Company", "Deputed Company", "Source", "Source Row", "Employee Code",
        "Employee Name", "UAN", "Aadhaar (masked)", "DOB", "DOJ", "Basic Wage",
        "Status", "Error Type", "Workflow Step", "Message", "Evidence", "Completed At",
    ]
    results_sheet.append(headers)
    for row in rows:
        results_sheet.append([row.get(header, "") for header in headers])
    events_sheet = workbook.create_sheet("Audit Events")
    events_sheet.append(["Timestamp", "Event", "Status", "Workflow Step", "Message"])
    for event in run.get("events", []):
        events_sheet.append([
            event.get("timestamp", ""), event.get("event", ""), event.get("status", ""),
            event.get("workflow_step", ""), event.get("message", ""),
        ])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column]
            sheet.column_dimensions[column[0].column_letter].width = min(max(max(values, default=10) + 2, 12), 55)
    xlsx_path = root / "processing_report.xlsx"
    xlsx_tmp = xlsx_path.with_suffix(".xlsx.tmp")
    workbook.save(xlsx_tmp)
    xlsx_tmp.replace(xlsx_path)
    return {
        "run_id": run_id,
        "status": run.get("status", ""),
        "generated_at": summary["generated_at"],
        "xlsx": {"name": xlsx_path.name, "size": xlsx_path.stat().st_size},
        "json": {"name": json_path.name, "size": json_path.stat().st_size},
        "summary": summary,
    }
